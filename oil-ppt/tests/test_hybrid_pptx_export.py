from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

import oil_ppt  # noqa: E402
import pptx_export  # noqa: E402


def make_png(path: Path, color: tuple[int, int, int], size: tuple[int, int] = (320, 180)) -> None:
    from PIL import Image
    Image.new("RGB", size, color).save(path)


class HybridPptxExportTests(unittest.TestCase):
    maxDiff = None

    def fixture(self, root: Path) -> tuple[dict, Path, Path]:
        project = root / "project"
        project.mkdir()
        (project / "assets").mkdir()
        media = project / "assets" / "evidence.png"
        make_png(media, (60, 120, 180), (400, 200))
        html = project / "演示文稿.html"
        html.write_text("<!doctype html><title>canonical</title>", encoding="utf-8")
        outline = {
            "title": "Hybrid deck",
            "slides": [
                {"id": "opening", "title": "Opening title", "image": "assets/evidence.png"},
                {"id": "detail", "title": "Detail title", "note": "A body note"},
            ],
        }
        return outline, project, html

    @staticmethod
    def fake_capture(
        _chrome: str, _html: Path, slides: list[dict], media: list[list[dict]], backgrounds: Path,
    ) -> list[dict]:
        layouts = []
        for index, slide in enumerate(slides):
            background = backgrounds / f"slide-{index + 1:03d}.png"
            make_png(background, (245 - index, 245, 240), (1920, 1080))
            text = [{
                "path": f"/slides/{index}/title", "text": slide["title"], "tag": "h1",
                "box": {"x": 80, "y": 70, "width": 800, "height": 100},
                "fontFamily": '"Noto Sans SC", sans-serif', "fontSize": "56px",
                "fontWeight": "700", "fontStyle": "normal", "color": "rgb(41, 41, 41)",
                "textAlign": "left", "lineHeight": "67px", "letterSpacing": "normal",
                "whiteSpace": "normal",
            }]
            if index == 1:
                text.append({
                    **text[0], "path": "/slides/1/note", "text": "A body note",
                    "box": {"x": 80, "y": 220, "width": 700, "height": 70},
                    "fontSize": "28px", "fontWeight": "400", "lineHeight": "36px",
                })
            images = []
            if media[index]:
                images.append({
                    "index": 0, "box": {"x": 1050, "y": 180, "width": 700, "height": 600},
                    "objectFit": "cover", "objectPosition": "right center", "filter": "none",
                    "borderRadius": "24px", "naturalWidth": 400, "naturalHeight": 200,
                })
            layouts.append({
                "slideId": slide["id"], "title": slide["title"], "text": text,
                "images": images, "fixedText": ["01"] if index == 0 else [],
                "background": str(background),
            })
        return layouts

    def test_export_writes_169_ordered_native_objects_and_consistent_report(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            outline, project, html = self.fixture(Path(temporary))
            output = project / "deck.pptx"
            report_path = project / "coverage.json"
            report = pptx_export.export_hybrid_pptx(
                project=project, outline=outline, html_path=html, chrome="unused",
                output=output, report_output=report_path, capture=self.fake_capture,
            )

            from pptx import Presentation
            presentation = Presentation(output)
            self.assertEqual(len(presentation.slides), 2)
            self.assertEqual(presentation.slide_width, pptx_export.SLIDE_WIDTH_EMU)
            self.assertEqual(presentation.slide_height, pptx_export.SLIDE_HEIGHT_EMU)
            self.assertEqual(
                [shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame],
                ["Opening title", "Detail title", "A body note"],
            )
            picture_counts = [
                sum(shape.shape_type == 13 for shape in slide.shapes)
                for slide in presentation.slides
            ]
            self.assertEqual(picture_counts, [2, 1])  # sanitized background + native project media
            native_picture = [shape for shape in presentation.slides[0].shapes if shape.shape_type == 13][1]
            self.assertGreater(native_picture.crop_left, 0)
            self.assertEqual(native_picture.crop_right, 0)
            self.assertEqual(native_picture._element.spPr.prstGeom.get("prst"), "roundRect")

            persisted = json.loads(report_path.read_text(encoding="utf-8"))
            self.assertEqual(persisted, report)
            self.assertEqual(report["summary"]["slide_count"], 2)
            self.assertEqual(report["summary"]["native_text_count"], 3)
            self.assertEqual(report["summary"]["native_media_count"], 1)
            self.assertEqual(report["slides"][0]["slide_id"], "opening")
            self.assertEqual(report["slides"][1]["slide_id"], "detail")
            self.assertEqual(report["slides"][0]["unsupported_structures"][0]["kind"], "unstructured-rendered-text")
            self.assertEqual(report["pptx_sha256"], pptx_export._sha256(output))

            with zipfile.ZipFile(output) as archive:
                names = archive.namelist()
                self.assertEqual(len([name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml")]), 2)
                self.assertGreaterEqual(len([name for name in names if name.startswith("ppt/media/")]), 3)
                for name in names:
                    if name.endswith(".rels"):
                        self.assertNotIn(b'TargetMode="External"', archive.read(name))
                theme = archive.read("ppt/theme/theme1.xml")
                self.assertIn(b"FFD54A", theme)
                self.assertIn(b"Noto Sans SC", theme)

    def test_atomic_failure_leaves_no_artifact_and_preserves_existing_pair(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            outline, project, html = self.fixture(Path(temporary))
            output = project / "deck.pptx"
            report_path = project / "coverage.json"
            output.write_bytes(b"old-pptx")
            report_path.write_text("old-report", encoding="utf-8")
            with mock.patch.object(pptx_export, "_validate_pptx", side_effect=RuntimeError("verification failed")):
                with self.assertRaisesRegex(RuntimeError, "verification failed"):
                    pptx_export.export_hybrid_pptx(
                        project=project, outline=outline, html_path=html, chrome="unused",
                        output=output, report_output=report_path, capture=self.fake_capture,
                    )
            self.assertEqual(output.read_bytes(), b"old-pptx")
            self.assertEqual(report_path.read_text(encoding="utf-8"), "old-report")
            self.assertFalse(list(project.glob(".deck.pptx.*.tmp")))
            self.assertFalse(list(project.glob(".coverage.json.*.tmp")))

    def test_dependency_failure_is_actionable(self) -> None:
        real_import = __import__

        def blocked_import(name: str, *args: object, **kwargs: object) -> object:
            if name == "pptx":
                raise ImportError("not installed")
            return real_import(name, *args, **kwargs)

        with mock.patch("builtins.__import__", side_effect=blocked_import):
            with self.assertRaisesRegex(SystemExit, "python -m pip install python-pptx"):
                pptx_export.require_python_pptx()

    def test_unsupported_svg_media_is_explicit_and_remains_rasterized(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            project = root / "project"
            project.mkdir()
            (project / "assets").mkdir()
            (project / "assets" / "diagram.svg").write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50"></svg>',
                encoding="utf-8",
            )
            html = project / "演示文稿.html"
            html.write_text("canonical", encoding="utf-8")
            outline = {"title": "SVG", "slides": [{"id": "one", "title": "One", "image": "assets/diagram.svg"}]}

            def capture(_chrome: str, _html: Path, slides: list[dict], _media: list[list[dict]], backgrounds: Path) -> list[dict]:
                background = backgrounds / "slide-001.png"
                make_png(background, (255, 255, 255), (1920, 1080))
                return [{"slideId": slides[0]["id"], "title": "One", "text": [], "images": [], "fixedText": [], "background": str(background)}]

            report = pptx_export.export_hybrid_pptx(
                project=project, outline=outline, html_path=html, chrome="unused",
                output=project / "deck.pptx", report_output=project / "coverage.json", capture=capture,
            )
            unsupported = report["slides"][0]["unsupported_structures"]
            self.assertEqual(unsupported[0]["kind"], "project-media")
            self.assertEqual(unsupported[0]["preserved_in"], "background")
            self.assertEqual(report["summary"]["native_media_count"], 0)


class PptxExportGateTests(unittest.TestCase):
    def test_command_is_discoverable_in_public_help(self) -> None:
        cli = SCRIPTS / "oil-ppt"
        result = subprocess.run(
            [str(cli), "--help"], capture_output=True, text=True, check=False, timeout=30,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("export-pptx", result.stdout)
        self.assertIn("coverage JSON", result.stdout)

    def test_stale_or_unconfirmed_project_is_rejected_before_export(self) -> None:
        project = Path("/tmp/oil-ppt-incomplete")
        with (
            mock.patch.object(oil_ppt, "require_initialized_project", return_value=project),
            mock.patch.object(oil_ppt, "require_no_edit_draft"),
            mock.patch.object(oil_ppt, "status_payload", return_value={
                "phase": "needs_preview_confirmation",
                "blockers": [{"message": "preview has not been confirmed"}],
            }),
        ):
            with self.assertRaisesRegex(SystemExit, "completed, confirmed, current"):
                oil_ppt.export_pptx_project(project)

    def test_active_editor_gate_runs_before_status_or_artifact_work(self) -> None:
        project = Path("/tmp/oil-ppt-active")
        with (
            mock.patch.object(oil_ppt, "require_initialized_project", return_value=project),
            mock.patch.object(oil_ppt, "require_no_edit_draft", side_effect=SystemExit("editor is open")),
            mock.patch.object(oil_ppt, "status_payload") as status,
        ):
            with self.assertRaisesRegex(SystemExit, "editor is open"):
                oil_ppt.export_pptx_project(project)
            status.assert_not_called()


if __name__ == "__main__":
    unittest.main()
