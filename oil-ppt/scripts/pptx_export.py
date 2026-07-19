#!/usr/bin/env python3
"""Hybrid PPTX export for a confirmed, browser-validated oil-ppt build.

The HTML deck remains canonical.  This module asks Chromium for the exact
rendered geometry, captures a background with supported structured objects
suppressed, and restores those objects as native PowerPoint shapes.
"""
from __future__ import annotations

import base64
import hashlib
import json
import os
import re
import shutil
import subprocess
import tempfile
import time
import urllib.request
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from cdp_validate import WebSocket, _stop_browser
from media_assets import inspect_image, outline_media_bindings
from palette_tokens import named_palette, normalize_palette
from profile_tokens import TYPE_PROFILES


SLIDE_WIDTH_PX = 1920
SLIDE_HEIGHT_PX = 1080
SLIDE_WIDTH_EMU = 12_192_000
SLIDE_HEIGHT_EMU = 6_858_000
COVERAGE_SCHEMA = "oil-ppt.pptx-editability/v1"


def require_python_pptx() -> Any:
    """Import python-pptx lazily so every non-export command remains stdlib-only."""
    try:
        import pptx
    except (ImportError, OSError) as error:
        raise SystemExit(
            "PPTX export requires the optional 'python-pptx' package. Install it for the "
            "same Python interpreter with: python -m pip install python-pptx"
        ) from error
    return pptx


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _cdp_command(socket: WebSocket, command_id: int, method: str, params: dict | None = None) -> dict:
    socket.send_json({"id": command_id, "method": method, **({"params": params} if params else {})})
    while True:
        message = socket.recv_json()
        if message.get("id") != command_id:
            continue
        if "error" in message:
            raise RuntimeError(f"Chrome DevTools {method} failed: {message['error']}")
        return message.get("result") or {}


def _browser_session(chrome: str, html_path: Path) -> tuple[subprocess.Popen, tempfile.TemporaryDirectory, WebSocket]:
    profile_ctx = tempfile.TemporaryDirectory(prefix="oil-ppt-pptx-cdp-")
    profile = Path(profile_ctx.name)
    command = [
        chrome, "--headless", "--no-sandbox", "--allow-file-access-from-files",
        "--disable-background-networking", "--disable-component-update", "--disable-default-apps",
        "--disable-features=PaintHolding,RenderDocument", "--disable-sync",
        "--force-color-profile=srgb", "--force-device-scale-factor=1", "--hide-scrollbars",
        "--metrics-recording-only", "--no-first-run", "--window-size=1920,1080",
        f"--user-data-dir={profile}", "--remote-debugging-port=0", html_path.resolve().as_uri(),
    ]
    process = subprocess.Popen(
        command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True,
    )
    try:
        deadline = time.monotonic() + 15
        port_file = profile / "DevToolsActivePort"
        while time.monotonic() < deadline and not port_file.exists():
            time.sleep(.05)
        if not port_file.exists():
            raise RuntimeError("Chrome did not expose a DevTools port for PPTX export.")
        port = int(port_file.read_text(encoding="utf-8").splitlines()[0])
        pages: list[dict] = []
        while time.monotonic() < deadline and not pages:
            try:
                with urllib.request.urlopen(f"http://127.0.0.1:{port}/json/list", timeout=2) as response:
                    pages = [item for item in json.load(response) if item.get("type") == "page"]
            except OSError:
                time.sleep(.05)
        if not pages:
            raise RuntimeError("Chrome opened no inspectable page for PPTX export.")
        target = html_path.resolve().as_uri()
        page = next((item for item in pages if item.get("url", "").startswith(target)), pages[0])
        socket = WebSocket(page["webSocketDebuggerUrl"])
        socket.sock.settimeout(20)
        return process, profile_ctx, socket
    except Exception:
        _stop_browser(process, profile)
        profile_ctx.cleanup()
        raise


def _layout_expression(slide_index: int, native_media_indexes: list[int]) -> str:
    """Return one self-contained browser expression for layout and suppression."""
    return f"""(async () => {{
      await document.fonts?.ready;
      await Promise.all([...document.images].map(image => image.complete
        ? Promise.resolve() : new Promise(resolve => {{ image.onload = image.onerror = resolve; }})));
      document.documentElement.style.cssText += ';width:1920px!important;height:1080px!important';
      document.body.style.cssText += ';width:1920px!important;height:1080px!important;overflow:hidden!important';
      const viewport = document.querySelector('.deck-viewport, .slide-preview-viewport');
      const shell = document.querySelector('.deck-stage-shell, .slide-preview-shell');
      const stage = document.querySelector('.deck-stage, .slide-preview-stage');
      if (!stage) throw new Error('canonical HTML has no deck stage');
      document.querySelectorAll('.deck-counter,.progress-bar,.next-preview').forEach(node => {{
        node.style.display = 'none';
      }});
      if (viewport) viewport.style.cssText += ';position:fixed!important;inset:0!important;width:1920px!important;height:1080px!important';
      if (shell) shell.style.cssText += ';position:relative!important;width:1920px!important;height:1080px!important';
      stage.style.cssText += ';position:absolute!important;left:0!important;top:0!important;width:1920px!important;height:1080px!important;transform:none!important';
      const slides = [...document.querySelectorAll('.oil-slide')];
      const slide = slides[{slide_index}];
      if (!slide) throw new Error('missing slide index {slide_index}');
      slides.forEach((node, index) => {{
        node.classList.toggle('active', index === {slide_index});
        node.style.cssText += index === {slide_index}
          ? ';display:block!important;visibility:visible!important;opacity:1!important;transition:none!important'
          : ';display:none!important;visibility:hidden!important;opacity:0!important;transition:none!important';
      }});
      await new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)));
      const visible = node => {{
        if (!node?.getClientRects().length) return false;
        const style = getComputedStyle(node);
        return style.display !== 'none' && style.visibility !== 'hidden' && Number(style.opacity || 1) > .001;
      }};
      const boxFor = node => {{
        const box = node.getBoundingClientRect();
        const root = slide.getBoundingClientRect();
        return {{x:box.left-root.left,y:box.top-root.top,width:box.width,height:box.height}};
      }};
      const text = [...slide.querySelectorAll('[data-edit-path]')]
        .filter(node => visible(node) && !node.querySelector('[data-edit-path]'))
        .map(node => {{
          const style = getComputedStyle(node);
          const value = (node.textContent || '').replace(/\\s+/g, ' ').trim();
          const color = style.color || '';
          const alpha = color.startsWith('rgba') ? Number(color.split(',').pop().replace(')','').trim()) : 1;
          const item = value && alpha > .001 ? {{
            path:node.dataset.editPath || '', text:value, tag:node.tagName.toLowerCase(),
            box:boxFor(node), fontFamily:style.fontFamily, fontSize:style.fontSize,
            fontWeight:style.fontWeight, fontStyle:style.fontStyle, color,
            textAlign:style.textAlign, lineHeight:style.lineHeight,
            letterSpacing:style.letterSpacing, whiteSpace:style.whiteSpace,
          }} : null;
          if (item) {{
            const walker = document.createTreeWalker(node, NodeFilter.SHOW_TEXT);
            const nodes = [];
            while (walker.nextNode()) nodes.push(walker.currentNode);
            nodes.forEach(textNode => {{
              const span = document.createElement('span');
              span.dataset.pptxHiddenText = 'true';
              span.style.visibility = 'hidden';
              textNode.parentNode.replaceChild(span, textNode);
              span.appendChild(textNode);
            }});
          }}
          return item;
        }}).filter(Boolean).filter(item => item.box.width > .5 && item.box.height > .5);
      const imageNodes = [...slide.querySelectorAll('img')].filter(visible);
      const images = imageNodes.map((node, index) => {{
        const style = getComputedStyle(node);
        return {{index,box:boxFor(node),objectFit:style.objectFit || 'fill',
          objectPosition:style.objectPosition || '50% 50%',filter:style.filter || 'none',
          borderRadius:style.borderRadius || '0px',naturalWidth:node.naturalWidth,
          naturalHeight:node.naturalHeight}};
      }});
      const fixedText = [];
      const walker = document.createTreeWalker(slide, NodeFilter.SHOW_TEXT);
      while (walker.nextNode()) {{
        const value = (walker.currentNode.nodeValue || '').replace(/\\s+/g, ' ').trim();
        const parent = walker.currentNode.parentElement;
        if (value && parent && visible(parent) && !parent.closest('[data-edit-path]')) fixedText.push(value);
      }}
      const nativeIndexes = new Set({json.dumps(native_media_indexes)});
      images.forEach(item => {{ if (nativeIndexes.has(item.index)) {{
        const node = imageNodes[item.index];
        if (node) node.style.visibility = 'hidden';
      }} }});
      await new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)));
      return {{slideId:slide.dataset.slideId || '', title:slide.dataset.title || '', text, images,
        fixedText:[...new Set(fixedText)], slideBox:boxFor(slide)}};
    }})()"""


def collect_render_layers(
    chrome: str,
    html_path: Path,
    slides: list[dict],
    media_by_slide: list[list[dict]],
    background_dir: Path,
) -> list[dict]:
    """Collect exact browser geometry and one sanitized PNG per slide."""
    process, profile_ctx, socket = _browser_session(chrome, html_path)
    profile = Path(profile_ctx.name)
    results: list[dict] = []
    command_id = 1
    try:
        _cdp_command(socket, command_id, "Page.enable"); command_id += 1
        _cdp_command(socket, command_id, "Runtime.enable"); command_id += 1
        for index, (slide, media) in enumerate(zip(slides, media_by_slide)):
            native_indexes = [item["rendered_index"] for item in media if item.get("native")]
            evaluated = _cdp_command(socket, command_id, "Runtime.evaluate", {
                "expression": _layout_expression(index, native_indexes),
                "awaitPromise": True,
                "returnByValue": True,
            }); command_id += 1
            remote = evaluated.get("result") or {}
            if remote.get("subtype") == "error" or "exceptionDetails" in evaluated:
                raise RuntimeError(f"Could not inspect rendered slide {index + 1}: {remote.get('description')}")
            layout = remote.get("value")
            if not isinstance(layout, dict):
                raise RuntimeError(f"Chrome returned no layout for slide {index + 1}.")
            if layout.get("slideId") != slide.get("id"):
                raise RuntimeError(
                    f"Rendered slide order mismatch at page {index + 1}: expected {slide.get('id')!r}, "
                    f"got {layout.get('slideId')!r}."
                )
            screenshot = _cdp_command(socket, command_id, "Page.captureScreenshot", {
                "format": "png", "fromSurface": True,
                "clip": {"x": 0, "y": 0, "width": SLIDE_WIDTH_PX, "height": SLIDE_HEIGHT_PX, "scale": 1},
            }); command_id += 1
            destination = background_dir / f"slide-{index + 1:03d}.png"
            destination.write_bytes(base64.b64decode(screenshot["data"]))
            layout["background"] = str(destination)
            results.append(layout)
    finally:
        socket.close()
        _stop_browser(process, profile)
        profile_ctx.cleanup()
    return results


def _media_plan(outline: dict, project: Path) -> list[list[dict]]:
    plans: list[list[dict]] = []
    for slide in outline["slides"]:
        items: list[dict] = []
        for rendered_index, (field, value) in enumerate(outline_media_bindings(slide)):
            path = (project / str(value)).resolve()
            try:
                info = inspect_image(path)
                native = info["mime"] in {"image/png", "image/jpeg", "image/gif", "image/webp"}
                reason = None if native else f"{info['mime']} cannot be embedded reliably by python-pptx"
            except ValueError as error:
                info = None
                native = False
                reason = str(error)
            items.append({
                "field": field, "source": str(value), "path": str(path),
                "rendered_index": rendered_index, "native": native, "reason": reason,
                **({"inspection": info} if info else {}),
            })
        plans.append(items)
    return plans


def _px_to_emu(value: float) -> int:
    return round(float(value) * SLIDE_WIDTH_EMU / SLIDE_WIDTH_PX)


def _rgb(value: str) -> tuple[int, int, int] | None:
    numbers = re.findall(r"[\d.]+", value or "")
    if len(numbers) < 3:
        return None
    return tuple(max(0, min(255, round(float(item)))) for item in numbers[:3])  # type: ignore[return-value]


def _font_name(value: str) -> str:
    first = (value or "Arial").split(",", 1)[0].strip().strip("\"'")
    return first or "Arial"


def _text_role(path: str) -> str:
    leaf = path.rsplit("/", 1)[-1]
    if leaf == "title":
        return "title"
    if leaf in {"body", "content", "statement_body", "render_body"}:
        return "body"
    if leaf in {"label", "kicker", "meta", "source", "caption", "page_note", "note"}:
        return "note-or-label"
    return "structured-copy"


def _position_fraction(token: str) -> float:
    token = token.strip().lower()
    if token in {"left", "top"}:
        return 0.0
    if token in {"right", "bottom"}:
        return 1.0
    if token == "center":
        return .5
    if token.endswith("%"):
        try:
            return max(0.0, min(1.0, float(token[:-1]) / 100))
        except ValueError:
            pass
    return .5


def _object_position(value: str) -> tuple[float, float]:
    tokens = value.split()
    if len(tokens) == 1:
        tokens *= 2
    return _position_fraction(tokens[0]), _position_fraction(tokens[1])


def _filtered_media(path: Path, css_filter: str, directory: Path, ordinal: int) -> tuple[Path, str | None]:
    if not css_filter or css_filter == "none":
        return path, None
    try:
        from PIL import Image, ImageEnhance, ImageOps
        image = Image.open(path).convert("RGBA")
        if "grayscale(1)" in css_filter:
            alpha = image.getchannel("A")
            image = ImageOps.grayscale(image.convert("RGB")).convert("RGBA")
            image.putalpha(alpha)
        saturation = re.search(r"saturate\(([\d.]+)\)", css_filter)
        contrast = re.search(r"contrast\(([\d.]+)\)", css_filter)
        if saturation:
            image = ImageEnhance.Color(image).enhance(float(saturation.group(1)))
        if contrast:
            image = ImageEnhance.Contrast(image).enhance(float(contrast.group(1)))
        destination = directory / f"filtered-{ordinal:03d}.png"
        image.save(destination, format="PNG")
        return destination, None
    except Exception as error:
        return path, f"CSS media filter could not be reproduced exactly: {error}"


def _add_picture(slide: Any, image_path: Path, rendered: dict, pptx: Any, temp_dir: Path, ordinal: int) -> str | None:
    box = rendered["box"]
    left, top = _px_to_emu(box["x"]), _px_to_emu(box["y"])
    width, height = _px_to_emu(box["width"]), _px_to_emu(box["height"])
    filtered, warning = _filtered_media(image_path, rendered.get("filter", "none"), temp_dir, ordinal)
    source_width = max(1, int(rendered.get("naturalWidth") or 1))
    source_height = max(1, int(rendered.get("naturalHeight") or 1))
    fit = rendered.get("objectFit") or "fill"
    x_fraction, y_fraction = _object_position(rendered.get("objectPosition") or "50% 50%")
    if fit == "contain":
        scale = min(width / source_width, height / source_height)
        picture_width, picture_height = round(source_width * scale), round(source_height * scale)
        left += round((width - picture_width) * x_fraction)
        top += round((height - picture_height) * y_fraction)
        picture = slide.shapes.add_picture(str(filtered), left, top, picture_width, picture_height)
    else:
        picture = slide.shapes.add_picture(str(filtered), left, top, width, height)
        if fit == "cover":
            source_ratio, box_ratio = source_width / source_height, width / height
            if source_ratio > box_ratio:
                visible = box_ratio / source_ratio
                picture.crop_left = (1 - visible) * x_fraction
                picture.crop_right = (1 - visible) * (1 - x_fraction)
            elif source_ratio < box_ratio:
                visible = source_ratio / box_ratio
                picture.crop_top = (1 - visible) * y_fraction
                picture.crop_bottom = (1 - visible) * (1 - y_fraction)
    if float(str(rendered.get("borderRadius") or "0").removesuffix("px") or 0) > 0:
        geometry = picture._element.spPr.prstGeom
        geometry.set("prst", "roundRect")
    return warning


def _set_east_asian_font(run: Any, typeface: str) -> None:
    from lxml import etree
    from pptx.oxml.ns import qn
    r_pr = run._r.get_or_add_rPr()
    east_asia = r_pr.find(qn("a:ea"))
    if east_asia is None:
        east_asia = etree.SubElement(r_pr, qn("a:ea"))
    east_asia.set("typeface", typeface)


def _add_text(slide: Any, item: dict, pptx: Any) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Pt
    box = item["box"]
    shape = slide.shapes.add_textbox(
        _px_to_emu(box["x"]), _px_to_emu(box["y"]),
        max(1, _px_to_emu(box["width"])), max(1, _px_to_emu(box["height"])),
    )
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = item.get("whiteSpace") not in {"nowrap", "pre"}
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = item["text"]
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.alignment = {
        "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }.get(item.get("textAlign"), PP_ALIGN.LEFT)
    run = paragraph.runs[0]
    font_name = _font_name(item.get("fontFamily", ""))
    run.font.name = font_name
    _set_east_asian_font(run, font_name)
    font_px = float(str(item.get("fontSize") or "16").removesuffix("px"))
    run.font.size = Pt(font_px * .75)
    weight = str(item.get("fontWeight") or "400")
    run.font.bold = weight == "bold" or (weight.isdigit() and int(weight) >= 600)
    run.font.italic = item.get("fontStyle") in {"italic", "oblique"}
    color = _rgb(item.get("color", ""))
    if color:
        run.font.color.rgb = RGBColor(*color)
    line_height = str(item.get("lineHeight") or "normal")
    if line_height.endswith("px"):
        paragraph.line_spacing = Pt(float(line_height[:-2]) * .75)


def _validate_pptx(path: Path, expected_slides: int, pptx: Any) -> None:
    reopened = pptx.Presentation(str(path))
    if len(reopened.slides) != expected_slides:
        raise RuntimeError(f"PPTX verification found {len(reopened.slides)} slides, expected {expected_slides}.")
    if reopened.slide_width != SLIDE_WIDTH_EMU or reopened.slide_height != SLIDE_HEIGHT_EMU:
        raise RuntimeError("PPTX verification found a non-16:9 slide size.")
    import zipfile
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.endswith(".rels"):
                continue
            if b'TargetMode="External"' in archive.read(name):
                raise RuntimeError(f"PPTX contains an external relationship in {name}.")


def _apply_ooxml_theme(path: Path, palette: dict[str, str], typeface: str) -> None:
    """Bind PowerPoint's editable-object defaults to the confirmed oil-ppt profile."""
    import xml.etree.ElementTree as ET
    import zipfile

    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ET.register_namespace("a", namespace)
    colors = {
        "dk1": palette["ink"], "lt1": palette["canvas"],
        "dk2": palette["ink_2"], "lt2": palette["surface"],
        "accent1": palette["accent"], "accent2": palette["accent_alt"],
        "accent3": palette["accent_warm"], "accent4": palette["accent_fill"],
        "accent5": palette["accent_soft"], "accent6": palette["border"],
    }
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".theme", dir=path.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "ppt/theme/theme1.xml":
                    root = ET.fromstring(content)
                    scheme = root.find(f".//{{{namespace}}}clrScheme")
                    if scheme is None:
                        raise RuntimeError("PPTX theme has no color scheme.")
                    for child in list(scheme):
                        name = child.tag.rsplit("}", 1)[-1]
                        if name not in colors:
                            continue
                        for existing in list(child):
                            child.remove(existing)
                        ET.SubElement(child, f"{{{namespace}}}srgbClr", {"val": colors[name].lstrip("#").upper()})
                    for font in root.findall(f".//{{{namespace}}}majorFont/{{{namespace}}}latin") + root.findall(
                        f".//{{{namespace}}}minorFont/{{{namespace}}}latin"
                    ):
                        font.set("typeface", typeface)
                    for font in root.findall(f".//{{{namespace}}}majorFont/{{{namespace}}}ea") + root.findall(
                        f".//{{{namespace}}}minorFont/{{{namespace}}}ea"
                    ):
                        font.set("typeface", typeface)
                    content = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target.writestr(info, content)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _install_pair_atomically(pptx_temp: Path, pptx_output: Path, report_temp: Path, report_output: Path) -> None:
    backups: list[tuple[Path, Path]] = []
    installed: list[Path] = []
    try:
        for destination in (pptx_output, report_output):
            if destination.exists():
                descriptor, backup_name = tempfile.mkstemp(
                    prefix=f".{destination.name}.", suffix=".bak", dir=destination.parent,
                )
                os.close(descriptor)
                Path(backup_name).unlink()
                backup = Path(backup_name)
                os.replace(destination, backup)
                backups.append((destination, backup))
        for temporary, destination in ((pptx_temp, pptx_output), (report_temp, report_output)):
            os.replace(temporary, destination)
            installed.append(destination)
    except Exception:
        for path in installed:
            path.unlink(missing_ok=True)
        for destination, backup in reversed(backups):
            os.replace(backup, destination)
        raise
    finally:
        for _, backup in backups:
            backup.unlink(missing_ok=True)


def export_hybrid_pptx(
    *,
    project: Path,
    outline: dict,
    html_path: Path,
    chrome: str,
    output: Path,
    report_output: Path,
    capture: Callable[[str, Path, list[dict], list[list[dict]], Path], list[dict]] = collect_render_layers,
) -> dict:
    """Build and atomically install a hybrid PPTX plus coverage report."""
    pptx = require_python_pptx()
    project, output, report_output = project.resolve(), output.resolve(), report_output.resolve()
    if output.suffix.lower() != ".pptx":
        raise SystemExit("PPTX output must use the .pptx extension.")
    if report_output.suffix.lower() != ".json":
        raise SystemExit("PPTX coverage report must use the .json extension.")
    if output == report_output:
        raise SystemExit("PPTX output and coverage report must be different files.")
    if not output.parent.is_dir() or not report_output.parent.is_dir():
        raise SystemExit("PPTX output and report parent directories must already exist.")
    media_by_slide = _media_plan(outline, project)
    temporary_root = Path(tempfile.mkdtemp(prefix=".oil-ppt-pptx-", dir=project))
    pptx_descriptor, pptx_name = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".tmp", dir=output.parent)
    os.close(pptx_descriptor)
    report_descriptor, report_name = tempfile.mkstemp(prefix=f".{report_output.name}.", suffix=".tmp", dir=report_output.parent)
    os.close(report_descriptor)
    pptx_temp, report_temp = Path(pptx_name), Path(report_name)
    try:
        backgrounds = temporary_root / "backgrounds"
        backgrounds.mkdir()
        layouts = capture(chrome, html_path, outline["slides"], media_by_slide, backgrounds)
        if len(layouts) != len(outline["slides"]):
            raise RuntimeError("Browser capture returned an incomplete slide set.")
        presentation = pptx.Presentation()
        presentation.slide_width = SLIDE_WIDTH_EMU
        presentation.slide_height = SLIDE_HEIGHT_EMU
        presentation.core_properties.title = str(outline.get("title") or "oil-ppt")
        presentation.core_properties.subject = "Hybrid editable export from canonical oil-ppt HTML"
        fixed_time = datetime(2000, 1, 1, tzinfo=timezone.utc)
        presentation.core_properties.created = fixed_time
        presentation.core_properties.modified = fixed_time
        while presentation.slides:
            slide_id = presentation.slides._sldIdLst[0]
            presentation.part.drop_rel(slide_id.rId)
            del presentation.slides._sldIdLst[0]
        per_slide: list[dict] = []
        total_text = total_media = total_unsupported = total_rasterized = 0
        for index, (outline_slide, media_plan, layout) in enumerate(
            zip(outline["slides"], media_by_slide, layouts), start=1,
        ):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(
                layout["background"], 0, 0, width=SLIDE_WIDTH_EMU, height=SLIDE_HEIGHT_EMU,
            )
            text_items = layout.get("text") or []
            for item in text_items:
                _add_text(slide, item, pptx)
            unsupported: list[dict] = []
            native_media = 0
            rendered_images = {int(item.get("index", -1)): item for item in layout.get("images") or []}
            for media_ordinal, planned in enumerate(media_plan, start=1):
                rendered = rendered_images.get(planned["rendered_index"])
                if not planned.get("native"):
                    unsupported.append({
                        "kind": "project-media", "field": planned["field"], "source": planned["source"],
                        "reason": planned.get("reason") or "unsupported media format",
                        "preserved_in": "background",
                    })
                    continue
                if rendered is None:
                    unsupported.append({
                        "kind": "project-media", "field": planned["field"], "source": planned["source"],
                        "reason": "structured media was not found in the rendered slide",
                        "preserved_in": "background",
                    })
                    continue
                warning = _add_picture(
                    slide, Path(planned["path"]), rendered, pptx, temporary_root, media_ordinal,
                )
                native_media += 1
                if warning:
                    unsupported.append({
                        "kind": "media-effect", "field": planned["field"], "source": planned["source"],
                        "reason": warning, "preserved_in": "native-media-with-approximation",
                    })
            fixed_text = layout.get("fixedText") or []
            if fixed_text:
                unsupported.append({
                    "kind": "unstructured-rendered-text", "count": len(fixed_text),
                    "examples": fixed_text[:8], "reason": "template-owned labels have no structured outline binding",
                    "preserved_in": "background",
                })
            rasterized_regions = [{
                "kind": "deterministic-background", "source": str(html_path),
                "sha256": _sha256(Path(layout["background"])),
                "excluded_native_text": len(text_items), "excluded_native_media": native_media,
                "contains": "CSS geometry, decorations, icons, and explicitly unsupported content",
            }]
            text_roles: dict[str, int] = {}
            for item in text_items:
                role = _text_role(str(item.get("path") or ""))
                text_roles[role] = text_roles.get(role, 0) + 1
            per_slide.append({
                "page": index, "slide_id": outline_slide["id"], "title": outline_slide["title"],
                "native_text_count": len(text_items), "native_text_by_role": text_roles,
                "native_media_count": native_media, "structured_media_count": len(media_plan),
                "unsupported_structure_count": sum(int(item.get("count", 1)) for item in unsupported),
                "rasterized_regions": rasterized_regions, "unsupported_structures": unsupported,
            })
            total_text += len(text_items)
            total_media += native_media
            total_unsupported += sum(int(item.get("count", 1)) for item in unsupported)
            total_rasterized += len(rasterized_regions)
        presentation.save(str(pptx_temp))
        raw_palette = outline.get("palette") or "oil-yellow"
        resolved_palette = named_palette(raw_palette) if isinstance(raw_palette, str) else normalize_palette(raw_palette)
        profile_name = str(outline.get("typography") or "clean")
        if profile_name not in TYPE_PROFILES:
            raise RuntimeError(f"Unknown typography profile during PPTX export: {profile_name}")
        theme_typeface = _font_name(TYPE_PROFILES[profile_name]["zh"])
        _apply_ooxml_theme(pptx_temp, resolved_palette, theme_typeface)
        with pptx_temp.open("rb") as stream:
            os.fsync(stream.fileno())
        _validate_pptx(pptx_temp, len(outline["slides"]), pptx)
        native_total = total_text + total_media
        eligible_total = native_total + total_unsupported
        report = {
            "schema_version": COVERAGE_SCHEMA, "ok": True,
            "project": str(project), "canonical_html": str(html_path),
            "pptx": str(output), "pptx_sha256": _sha256(pptx_temp),
            "slide_size": {"width_emu": SLIDE_WIDTH_EMU, "height_emu": SLIDE_HEIGHT_EMU, "aspect_ratio": "16:9"},
            "theme": {"palette": resolved_palette, "typography": profile_name, "typeface": theme_typeface},
            "summary": {
                "slide_count": len(per_slide), "native_text_count": total_text,
                "native_media_count": total_media, "native_object_count": native_total,
                "rasterized_region_count": total_rasterized,
                "unsupported_structure_count": total_unsupported,
                "coverage_denominator": eligible_total,
                "editable_coverage_percent": round(100 * native_total / eligible_total, 2) if eligible_total else 100.0,
            },
            "methodology": (
                "Coverage counts native structured text/media objects against those objects plus explicitly "
                "unsupported structures. CSS geometry and decoration are preserved in one sanitized raster background per slide."
            ),
            "slides": per_slide,
        }
        with report_temp.open("w", encoding="utf-8") as stream:
            json.dump(report, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        _install_pair_atomically(pptx_temp, output, report_temp, report_output)
        return report
    finally:
        pptx_temp.unlink(missing_ok=True)
        report_temp.unlink(missing_ok=True)
        shutil.rmtree(temporary_root, ignore_errors=True)
